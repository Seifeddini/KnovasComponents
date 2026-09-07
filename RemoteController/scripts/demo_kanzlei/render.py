"""Render corpus documents to the four ingestible formats.

Long documents: ``.txt``, ``.docx``, ``.pdf`` (text-layer).
Emails: Outlook ``.msg`` (OLE compound file), not ``.eml``.
Scanned incoming post: image-only PDF (separate helper).
"""
from __future__ import annotations

import html
import textwrap
from datetime import datetime
from pathlib import Path

import fitz
from docx import Document
from docx.shared import Pt
from msgforge import Message
from PIL import Image, ImageDraw, ImageFont

SYNCABLE_LONG_FORMATS = ("txt", "docx", "pdf")
EMAIL_FORMAT = "msg"


def render_txt(path: Path, body: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(body, encoding="utf-8")


def render_docx(path: Path, body: str, title: str | None = None) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    if title:
        heading = doc.add_heading(title, level=1)
        for run in heading.runs:
            run.font.size = Pt(14)
    for paragraph in body.split("\n"):
        doc.add_paragraph(paragraph)
    doc.save(str(path))


def render_pdf(path: Path, body: str, title: str | None = None) -> None:
    """Text-layer PDF so extraction works without OCR."""
    path.parent.mkdir(parents=True, exist_ok=True)
    doc = fitz.open()
    page = doc.new_page()
    rect = fitz.Rect(54, 54, page.rect.width - 54, page.rect.height - 54)
    escaped = html.escape(body).replace("\n", "<br/>")
    heading = f"<h1>{html.escape(title)}</h1>" if title else ""
    page.insert_htmlbox(rect, f"{heading}<p style='font-size:11pt;line-height:1.35'>{escaped}</p>")
    doc.save(str(path), deflate=True)
    doc.close()


def render_msg(
    path: Path,
    *,
    subject: str,
    body: str,
    sender: tuple[str, str],
    to: list[tuple[str, str]],
    sent: datetime,
    cc: list[tuple[str, str]] | None = None,
) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    message = Message(
        subject=subject,
        text_body=body,
        sender=sender,
        to=to,
        cc=cc or [],
        sent=sent,
    )
    message.save(str(path))


def render_scan_pdf(path: Path, body: str) -> None:
    """Image-only PDF (no text layer) for the OCR ingest path."""
    path.parent.mkdir(parents=True, exist_ok=True)
    width, height = 1240, 1754  # ~150 dpi A4
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    try:
        font = ImageFont.truetype("arial.ttf", 28)
    except OSError:
        font = ImageFont.load_default()
    y = 60
    for line in _wrap_for_scan(body, 72):
        draw.text((60, y), line, fill="black", font=font)
        y += 36
        if y > height - 80:
            break
    pdf = fitz.open()
    page = pdf.new_page(width=595, height=842)
    buf = _png_bytes(image)
    page.insert_image(page.rect, stream=buf)
    pdf.save(str(path), deflate=True)
    pdf.close()


def _png_bytes(image: Image.Image) -> bytes:
    from io import BytesIO

    buf = BytesIO()
    image.save(buf, format="PNG")
    return buf.getvalue()


def _wrap_for_scan(body: str, width: int) -> list[str]:
    lines: list[str] = []
    for raw in body.splitlines() or [""]:
        if not raw.strip():
            lines.append("")
            continue
        lines.extend(textwrap.wrap(raw, width=width) or [""])
    return lines


def render_document(
    path: Path,
    body: str,
    *,
    fmt: str,
    title: str | None = None,
    subject: str | None = None,
    sender: tuple[str, str] | None = None,
    to: list[tuple[str, str]] | None = None,
    sent: datetime | None = None,
) -> None:
    kind = fmt.lstrip(".").lower()
    if kind == "txt":
        render_txt(path, body)
    elif kind == "docx":
        render_docx(path, body, title=title)
    elif kind == "pdf":
        render_pdf(path, body, title=title)
    elif kind == "scan_pdf":
        render_scan_pdf(path, body)
    elif kind == "msg":
        if not (subject and sender and to and sent):
            raise ValueError("msg render requires subject, sender, to, sent")
        render_msg(path, subject=subject, body=body, sender=sender, to=to, sent=sent)
    else:
        raise ValueError(f"unsupported render format: {fmt}")
